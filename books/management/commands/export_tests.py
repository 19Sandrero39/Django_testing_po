import inspect
import textwrap
import unittest
from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand
from django.test.runner import DiscoverRunner
from docx import Document
from docx.shared import Pt


def iter_tests(suite):
    """Разворачивает вложенные наборы тестов в плоский список."""
    for item in suite:
        if isinstance(item, unittest.TestSuite):
            yield from iter_tests(item)
        else:
            yield item


def describe(test):
    """Возвращает (класс, имя, описание, исходный код) теста."""
    cls = type(test)
    method = getattr(cls, test._testMethodName)
    source = textwrap.dedent(inspect.getsource(method))
    description = inspect.getdoc(method)
    if not description:
        # Берём первый комментарий внутри теста как описание
        for line in source.splitlines()[1:]:
            stripped = line.strip()
            if stripped.startswith("#"):
                description = stripped.lstrip("# ").strip()
                break
            if stripped:
                break
    return cls.__name__, test._testMethodName, description or "—", source


class Command(BaseCommand):
    help = "Записывает все тесты проекта в Word-файл"

    def add_arguments(self, parser):
        parser.add_argument(
            "--output", default="tests_report.docx", help="имя итогового файла"
        )
        parser.add_argument(
            "--run", action="store_true", help="запустить тесты и добавить результат"
        )

    def handle(self, *args, **options):
        runner = DiscoverRunner(verbosity=0)
        suite = runner.build_suite()
        tests = list(iter_tests(suite))
        items = [(test.id(),) + describe(test) for test in tests]

        failed = set()
        if options["run"]:
            runner.setup_test_environment()
            old_config = runner.setup_databases()
            try:
                result = unittest.TestResult()
                suite.run(result)
            finally:
                runner.teardown_databases(old_config)
                runner.teardown_test_environment()
            failed = {t.id() for t, _ in result.failures + result.errors}

        doc = Document()
        doc.add_heading("Тесты Django-проекта", 0)
        doc.add_paragraph(
            f"Дата: {datetime.now():%d.%m.%Y %H:%M}. Всего тестов: {len(items)}."
        )

        # Сводная таблица
        headers = ["№", "Класс", "Тест", "Описание"]
        if options["run"]:
            headers.append("Результат")
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for number, (test_id, cls, name, description, _) in enumerate(items, 1):
            row = [str(number), cls, name, description]
            if options["run"]:
                row.append("FAIL" if test_id in failed else "OK")
            for cell, text in zip(table.add_row().cells, row):
                cell.text = text

        # Исходный код каждого теста
        doc.add_heading("Код тестов", 1)
        current_class = None
        for _, cls, name, description, source in items:
            if cls != current_class:
                doc.add_heading(cls, 2)
                current_class = cls
            doc.add_paragraph(f"{name}: {description}")
            code = doc.add_paragraph().add_run(source)
            code.font.name = "Consolas"
            code.font.size = Pt(9)

        path = Path(options["output"])
        doc.save(path)
        self.stdout.write(self.style.SUCCESS(f"Готово: {path.resolve()}"))
